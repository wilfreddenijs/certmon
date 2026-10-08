import io
import json
import sqlite3

import pytest
from openpyxl import load_workbook

from certmon.auth import hash_password
from tests.test_auth_api import load_app
from tests.test_audit_api import auth_headers


def seed_events(module):
    with module.database.transaction() as connection:
        connection.executemany(
            """INSERT INTO audit_events(event_type, target, success, details_json, created_at)
               VALUES (?, ?, ?, ?, ?)""",
            [("certificate_upload_succeeded", "10.0.0.1", 1,
              json.dumps({"certificate_id": "cert-1"}), "2026-01-01T23:59:59+00:00"),
             ("certificate_upload_failed", "10.0.0.2", 0, "{}", "2026-01-02T00:00:00+00:00"),
             ("=HYPERLINK(\"https://example.com\")", "=1+1", 1, "{}", "2026-01-03T00:00:00+00:00")],
        )


def test_excel_exports_entire_audit_and_keeps_formulas_as_text(tmp_data_dir, monkeypatch):
    module = load_app(tmp_data_dir, monkeypatch, server_mode=False)
    seed_events(module)
    for index in range(505):
        module.audit_service.record("extra", target=str(index))
    response = module.app.test_client().get("/api/audit/export/excel")
    assert response.status_code == 200
    assert "attachment" in response.headers["Content-Disposition"]
    workbook = load_workbook(io.BytesIO(response.data))
    sheet = workbook["Audit"]
    assert sheet.max_row == 509
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert any(row[4:6] == ("10.0.0.2", "Failed") for row in rows)
    assert any(row[4] == "10.0.0.1" and "cert-1" in row[6] for row in rows)
    assert all(cell.data_type != "f" for row in sheet for cell in row)
    assert any(row[4] == "=1+1" for row in rows)
    assert sheet.freeze_panes == "A2"


def test_cleanup_preserves_date_boundary_and_records_deletion(tmp_data_dir, monkeypatch):
    module = load_app(tmp_data_dir, monkeypatch, server_mode=False)
    seed_events(module)
    client = module.app.test_client()
    response = client.post("/api/audit/cleanup", json={"before": "2026-01-02", "confirm": True})
    assert response.status_code == 200
    assert response.get_json()["deleted_count"] == 1
    events = module.database.list_audit_events()
    assert len(events) == 3
    assert events[0]["event_type"] == "audit_history_deleted"
    assert events[0]["details"]["deleted_count"] == 1
    assert any(event["target"] == "10.0.0.2" for event in events)
    assert not any(event["target"] == "10.0.0.1" for event in events)


@pytest.mark.parametrize("body", [None, [], {}, {"before": "2026-01-02"},
    {"before": "2026-02-30", "confirm": True}, {"before": "2099-01-01", "confirm": True},
    {"before": "2026-1-2", "confirm": True}, {"before": None, "confirm": True}])
def test_invalid_cleanup_never_changes_history(tmp_data_dir, monkeypatch, body):
    module = load_app(tmp_data_dir, monkeypatch, server_mode=False)
    seed_events(module)
    before = module.database.list_audit_events()
    response = module.app.test_client().post("/api/audit/cleanup", json=body)
    assert response.status_code == 400
    assert module.database.list_audit_events() == before


@pytest.mark.parametrize("role,export_status,delete_status", [
    ("viewer", 403, 403), ("operator", 403, 403), ("security_admin", 200, 403), ("admin", 200, 200)])
def test_export_and_cleanup_roles(tmp_data_dir, monkeypatch, role, export_status, delete_status):
    module = load_app(tmp_data_dir, monkeypatch)
    module.database.create_user(user_id="test-user", username="test-user",
                                password_hash=hash_password("correct horse"), roles=[role])
    client = module.app.test_client()
    client.post("/api/auth/login", json={"username": "test-user", "password": "correct horse"})
    assert client.get("/api/audit/export/excel").status_code == export_status
    body = {"before": "2026-01-02", "confirm": True}
    assert client.post("/api/audit/cleanup", json=body).status_code == 403
    assert client.post("/api/audit/cleanup", json=body, headers=auth_headers(client)).status_code == delete_status


def test_audit_actions_require_login(tmp_data_dir, monkeypatch):
    module = load_app(tmp_data_dir, monkeypatch)
    client = module.app.test_client()
    assert client.get("/api/audit/export/excel").status_code == 401
    assert client.post("/api/audit/cleanup", json={"before": "2026-01-02", "confirm": True}).status_code == 401


def test_deletion_rolls_back_if_its_audit_record_cannot_be_saved(tmp_data_dir, monkeypatch):
    module = load_app(tmp_data_dir, monkeypatch, server_mode=False)
    seed_events(module)
    before = module.database.list_audit_events()
    with module.database.transaction() as connection:
        connection.execute("""CREATE TRIGGER reject_cleanup_audit BEFORE INSERT ON audit_events
            WHEN NEW.event_type = 'audit_history_deleted'
            BEGIN SELECT RAISE(ABORT, 'cannot save audit'); END""")
    with pytest.raises(sqlite3.IntegrityError):
        module.database.delete_audit_before("2026-01-02T00:00:00+00:00")
    assert module.database.list_audit_events() == before
